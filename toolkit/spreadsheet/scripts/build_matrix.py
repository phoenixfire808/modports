from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHEETS = {
    "Entities": ["entity_id", "name", "entity_type", "enabled", "description"],
    "Fields": ["field_id", "entity_id", "value_type", "value", "unit"],
    "Rules": ["rule_id", "entity_id", "rule_type", "target", "value", "enabled"],
}
CODE = re.compile(r"^[A-Z][A-Z0-9_]{2,63}$")
ENTITY_TYPES = {"item", "weapon", "npc", "resource", "other"}
VALUE_TYPES = {"string", "integer", "number", "boolean", "entity_ref"}
RULE_TYPES = {"set_field", "requires_entity", "cooldown"}


class ValidationError(Exception):
    pass


def scalar(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def rows_for(workbook: Any, sheet_name: str, headers: list[str], errors: list[str]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        errors.append(f"Missing required sheet: {sheet_name}")
        return []
    sheet = workbook[sheet_name]
    for cells in sheet.iter_rows():
        for cell in cells:
            if cell.data_type in {"f", "e"}:
                errors.append(f"{sheet_name}!{cell.coordinate}: formulas and Excel errors are not accepted; use a literal value")
    actual = [cell.value for cell in sheet[1]]
    if actual != headers:
        errors.append(f"{sheet_name}!1 headers must be exactly: {', '.join(headers)}")
        return []
    result = []
    for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        data = dict(zip(headers, values))
        if all(value is None or value == "" for value in data.values()):
            continue
        result.append({key: scalar(value) for key, value in data.items()})
    return result


def require_text(record: dict[str, Any], key: str, label: str, errors: list[str]) -> str:
    value = record.get(key)
    if not isinstance(value, str) or not value.strip():
        errors.append(f"{label}: {key} is required and must be text")
        return ""
    return value.strip()


def validate(workbook: Any) -> dict[str, list[dict[str, Any]]]:
    errors: list[str] = []
    data = {name: rows_for(workbook, name, headers, errors) for name, headers in SHEETS.items()}
    entities = data["Entities"]
    fields = data["Fields"]
    rules = data["Rules"]

    def unique_codes(records: list[dict[str, Any]], key: str, sheet: str) -> set[str]:
        found: set[str] = set()
        for index, record in enumerate(records, start=2):
            code = require_text(record, key, f"{sheet} row {index}", errors)
            if code and not CODE.fullmatch(code):
                errors.append(f"{sheet} row {index}: invalid {key} {code!r}; expected uppercase stable code")
            if code in found:
                errors.append(f"{sheet} row {index}: duplicate {key} {code!r}")
            found.add(code)
        return found

    entity_ids = unique_codes(entities, "entity_id", "Entities")
    field_ids = unique_codes(fields, "field_id", "Fields")
    unique_codes(rules, "rule_id", "Rules")

    for index, row in enumerate(entities, start=2):
        require_text(row, "name", f"Entities row {index}", errors)
        typ = require_text(row, "entity_type", f"Entities row {index}", errors)
        if typ and typ not in ENTITY_TYPES:
            errors.append(f"Entities row {index}: unsupported entity_type {typ!r}")
        enabled = row.get("enabled")
        if not isinstance(enabled, bool):
            errors.append(f"Entities row {index}: enabled must be TRUE/FALSE")

    for index, row in enumerate(fields, start=2):
        owner = require_text(row, "entity_id", f"Fields row {index}", errors)
        if owner and owner not in entity_ids:
            errors.append(f"Fields row {index}: unknown entity_id {owner!r}")
        typ = require_text(row, "value_type", f"Fields row {index}", errors)
        value = row.get("value")
        if typ and typ not in VALUE_TYPES:
            errors.append(f"Fields row {index}: unsupported value_type {typ!r}")
        elif typ and value is None:
            errors.append(f"Fields row {index}: value is required")
        elif typ == "integer" and (isinstance(value, bool) or not isinstance(value, int)):
            errors.append(f"Fields row {index}: integer value required")
        elif typ == "number" and (isinstance(value, bool) or not isinstance(value, (int, float))):
            errors.append(f"Fields row {index}: numeric value required")
        elif typ == "boolean" and not isinstance(value, bool):
            errors.append(f"Fields row {index}: boolean value required")
        elif typ in {"string", "entity_ref"} and not isinstance(value, str):
            errors.append(f"Fields row {index}: text value required")
        elif typ == "entity_ref" and value not in entity_ids:
            errors.append(f"Fields row {index}: entity_ref {value!r} does not exist")

    for index, row in enumerate(rules, start=2):
        owner = require_text(row, "entity_id", f"Rules row {index}", errors)
        if owner and owner not in entity_ids:
            errors.append(f"Rules row {index}: unknown entity_id {owner!r}")
        typ = require_text(row, "rule_type", f"Rules row {index}", errors)
        if typ and typ not in RULE_TYPES:
            errors.append(f"Rules row {index}: unsupported rule_type {typ!r}")
        target = require_text(row, "target", f"Rules row {index}", errors)
        if typ == "set_field" and target not in field_ids:
            errors.append(f"Rules row {index}: set_field target {target!r} is not a field_id")
        if typ == "requires_entity" and target not in entity_ids:
            errors.append(f"Rules row {index}: requires_entity target {target!r} is not an entity_id")
        if typ == "cooldown" and target not in field_ids:
            errors.append(f"Rules row {index}: cooldown target {target!r} is not a field_id")
        if not isinstance(row.get("enabled"), bool):
            errors.append(f"Rules row {index}: enabled must be TRUE/FALSE")
        if row.get("value") is None:
            errors.append(f"Rules row {index}: value is required (use a literal such as TRUE for no parameter)")

    if errors:
        raise ValidationError("\n".join(errors))
    return data


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate XLSX source workbook and generate canonical model matrix JSON")
    parser.add_argument("--workbook", type=Path, default=ROOT / "workbooks" / "model.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "generated" / "matrix.json")
    parser.add_argument("--check", action="store_true", help="Fail on missing/stale output without writing")
    args = parser.parse_args()
    if args.workbook.resolve() == args.output.resolve():
        print("ERROR: source workbook and output must be different files", file=sys.stderr)
        return 2
    if not args.workbook.is_file():
        print(f"ERROR: workbook not found: {args.workbook}", file=sys.stderr)
        return 2
    try:
        workbook = load_workbook(args.workbook, data_only=False, read_only=True, keep_links=False)
        try:
            data = validate(workbook)
        finally:
            workbook.close()
    except ValidationError as exc:
        print(f"Workbook validation failed:\n{exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"Could not read workbook: {exc}", file=sys.stderr)
        return 2

    digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    payload = {"schema_version": 1, "source": {"workbook": args.workbook.name, "sha256": digest}, "tables": data}
    rendered = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.check:
        if not args.output.is_file() or args.output.read_bytes() != rendered.encode("utf-8"):
            print("STALE: run build_matrix.py without --check", file=sys.stderr)
            return 1
        print("Matrix is current")
        return 0
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=args.output.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(rendered.encode("utf-8"))
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, args.output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    print(f"Valid: {sum(len(rows) for rows in data.values())} rows. Wrote {args.output}")
    print(f"Source SHA-256: {digest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
