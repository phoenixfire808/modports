from pathlib import Path
from openpyxl import Workbook

root = Path(__file__).resolve().parents[1]
path = root / "workbooks" / "model.xlsx"
if path.exists():
    raise SystemExit("Refusing to overwrite an existing workbook. Use a fresh project directory.")
path.parent.mkdir(parents=True, exist_ok=True)
book = Workbook()
entities = book.active
entities.title = "Entities"
entities.append(["entity_id", "name", "entity_type", "enabled", "description"])
entities.append(["ENT_WOOD", "Wood", "resource", True, "Illustrative template row, replace with approved content."])
entities.append(["ENT_HATCHET", "Hatchet", "weapon", True, "Illustrative template row, replace with approved content."])
fields = book.create_sheet("Fields")
fields.append(["field_id", "entity_id", "value_type", "value", "unit"])
fields.append(["FLD_WOOD_STACK", "ENT_WOOD", "integer", 100, "items"])
fields.append(["FLD_HATCHET_COOLDOWN", "ENT_HATCHET", "number", 1.2, "seconds"])
rules = book.create_sheet("Rules")
rules.append(["rule_id", "entity_id", "rule_type", "target", "value", "enabled"])
rules.append(["RUL_HATCHET_COOLDOWN", "ENT_HATCHET", "cooldown", "FLD_HATCHET_COOLDOWN", 1.2, True])
evidence = book.create_sheet("Evidence")
evidence.append(["evidence_id", "target_id", "status", "observation", "source", "version", "procedure", "expected", "observed", "test_id"])
evidence.append(["EVD_DEMO", "FLD_HATCHET_COOLDOWN", "illustrative", "Made-up starter value, not recovered behavior", "Original synthetic example", "DEMO_1", "Replace with a repeatable measurement", "Unknown", "Not measured", "TST_DEMO"])
for sheet in book.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 60)
book.save(path)
print(f"Created {path}")
