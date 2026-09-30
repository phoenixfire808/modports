import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from build_matrix import ValidationError, validate


class WorkbookValidationTests(unittest.TestCase):
    def workbook(self):
        book = Workbook()
        entities = book.active
        entities.title = "Entities"
        entities.append(["entity_id", "name", "entity_type", "enabled", "description"])
        entities.append(["ENT_WOOD", "Wood", "resource", True, "sample"])
        fields = book.create_sheet("Fields")
        fields.append(["field_id", "entity_id", "value_type", "value", "unit"])
        fields.append(["FLD_STACK", "ENT_WOOD", "integer", 50, "items"])
        rules = book.create_sheet("Rules")
        rules.append(["rule_id", "entity_id", "rule_type", "target", "value", "enabled"])
        rules.append(["RUL_STACK", "ENT_WOOD", "set_field", "FLD_STACK", 50, True])
        return book

    def test_valid_workbook_normalizes_all_tables(self):
        data = validate(self.workbook())
        self.assertEqual(["ENT_WOOD"], [row["entity_id"] for row in data["Entities"]])
        self.assertEqual(50, data["Fields"][0]["value"])
        self.assertEqual("FLD_STACK", data["Rules"][0]["target"])

    def test_duplicate_entity_code_is_rejected(self):
        book = self.workbook()
        book["Entities"].append(["ENT_WOOD", "Duplicate", "resource", True, ""])
        with self.assertRaisesRegex(ValidationError, "duplicate entity_id"):
            validate(book)

    def test_broken_entity_reference_is_rejected(self):
        book = self.workbook()
        book["Fields"]["B2"] = "ENT_MISSING"
        with self.assertRaisesRegex(ValidationError, "unknown entity_id"):
            validate(book)

    def test_rule_target_must_reference_known_field(self):
        book = self.workbook()
        book["Rules"]["D2"] = "FLD_MISSING"
        with self.assertRaisesRegex(ValidationError, "set_field target"):
            validate(book)


if __name__ == "__main__":
    unittest.main()
