import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

class StarterCliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'scripts').mkdir()
        shutil.copyfile(ROOT / 'scripts/create_template.py', self.root / 'scripts/create_template.py')
        self.workbook = self.root / 'workbooks/model.xlsx'
        self.output = self.root / 'generated/matrix.json'
        result = self.generate()
        self.assertEqual(result.returncode, 0, result.stderr)

    def generate(self):
        return subprocess.run([sys.executable, str(self.root / 'scripts/create_template.py')], capture_output=True, text=True)

    def run_cli(self, *extra, output=None):
        return subprocess.run([sys.executable, str(ROOT / 'scripts/build_matrix.py'), '--workbook', str(self.workbook), '--output', str(output or self.output), *extra], capture_output=True, text=True)

    def edit(self, sheet, cell, value):
        book = load_workbook(self.workbook)
        book[sheet][cell] = value
        book.save(self.workbook)
        book.close()

    def test_build_and_check_are_deterministic(self):
        self.assertEqual(self.run_cli().returncode, 0)
        original = self.output.read_bytes()
        self.assertEqual(self.run_cli().returncode, 0)
        self.assertEqual(self.output.read_bytes(), original)
        self.assertEqual(self.run_cli('--check').returncode, 0)
        data = json.loads(original)
        self.assertEqual(data['tables']['Fields'][1]['value'], 1.2)
        self.assertNotIn(str(self.root), original.decode())

    def test_missing_output_check_does_not_write(self):
        self.assertEqual(self.run_cli('--check').returncode, 1)
        self.assertFalse(self.output.exists())

    def test_stale_check_preserves_last_output(self):
        self.assertEqual(self.run_cli().returncode, 0)
        original = self.output.read_bytes()
        self.edit('Fields', 'D2', 101)
        self.assertEqual(self.run_cli('--check').returncode, 1)
        self.assertEqual(self.output.read_bytes(), original)

    def test_formula_rejected_without_replacing_output(self):
        self.assertEqual(self.run_cli().returncode, 0)
        original = self.output.read_bytes()
        self.edit('Entities', 'B2', '=1+1')
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        self.assertIn('formulas', result.stderr)
        self.assertEqual(self.output.read_bytes(), original)

    def test_excel_error_is_rejected(self):
        self.edit('Entities', 'B2', '#REF!')
        self.assertEqual(self.run_cli().returncode, 1)
        self.assertFalse(self.output.exists())

    def test_source_cannot_be_output(self):
        original = self.workbook.read_bytes()
        self.assertEqual(self.run_cli(output=self.workbook).returncode, 2)
        self.assertEqual(self.workbook.read_bytes(), original)

    def test_generator_never_overwrites_existing_workbook(self):
        original = self.workbook.read_bytes()
        self.assertNotEqual(self.generate().returncode, 0)
        self.assertEqual(self.workbook.read_bytes(), original)

    def test_evidence_is_explicitly_not_exported(self):
        book = load_workbook(self.workbook)
        self.assertIn('Evidence', book.sheetnames)
        book.close()
        self.assertEqual(self.run_cli().returncode, 0)
        self.assertEqual(set(json.loads(self.output.read_bytes())['tables']), {'Entities', 'Fields', 'Rules'})

if __name__ == '__main__':
    unittest.main()
